from io import BytesIO

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

HEADER_FILL = PatternFill("solid", fgColor="3171F0")
HEADER_FONT = Font(color="FFFFFF", bold=True)


TEMPLATES = {
    "people": ["name", "email", "phone", "job_title", "organization", "owner_email"],
    "organizations": ["name", "address", "website", "phone", "owner_email"],
    "deals": [
        "title",
        "value",
        "currency",
        "stage",
        "status",
        "organization",
        "contact_person",
        "owner_email",
        "expected_close_date",
    ],
    "leads": [
        "title",
        "person_name",
        "email",
        "phone",
        "organization_name",
        "source",
        "status",
        "owner_email",
    ],
    "activities": [
        "type",
        "subject",
        "due_at",
        "done",
        "person",
        "organization",
        "deal",
        "owner_email",
    ],
    "customers": [
        "customer_name",
        "customer_email",
        "customer_phone",
        "organization_name",
        "request_type",
        "title",
        "description",
        "status",
    ],
}


def workbook_from_rows(headers, rows, sheet_name="Export"):
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    for row in rows:
        ws.append(list(row))
    for column in ws.columns:
        max_length = 12
        col_letter = column[0].column_letter
        for cell in column:
            if cell.value:
                max_length = max(max_length, min(len(str(cell.value)), 40))
        ws.column_dimensions[col_letter].width = max_length + 2
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def template_bytes(kind):
    headers = TEMPLATES[kind]
    return workbook_from_rows(headers, [], sheet_name="Template")


def read_tabular(file_storage):
    filename = (file_storage.filename or "").lower()
    if filename.endswith(".csv"):
        import csv
        import io

        text = file_storage.stream.read().decode("utf-8-sig")
        reader = csv.DictReader(io.StringIO(text))
        return reader.fieldnames or [], list(reader)
    wb = load_workbook(file_storage, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return [], []
    headers = [str(h).strip() if h is not None else "" for h in rows[0]]
    records = []
    for row in rows[1:]:
        if not row or all(v is None or str(v).strip() == "" for v in row):
            continue
        records.append(
            {
                headers[i]: ("" if row[i] is None else str(row[i]).strip())
                for i in range(len(headers))
                if headers[i]
            }
        )
    return headers, records


def mapped_record(raw, mapping):
    result = {}
    for field, source in mapping.items():
        if not source or source == "__skip__":
            continue
        result[field] = (raw.get(source) or "").strip()
    return result
